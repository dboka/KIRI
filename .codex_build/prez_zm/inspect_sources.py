from pathlib import Path
import json
import zipfile

from docx import Document
from openpyxl import load_workbook

root = Path(r"C:/Users/deniss.boka/MESLI_PROJECT/KIRI")
workbook_path = Path(r"C:/Users/deniss.boka/Downloads/Darbu_plans.xlsx")
docx_path = root / "docs" / "KIRI_LV_pilns_plans_un_arhitektura.docx"
pptm_path = root / "PREZ.pptm"
out_dir = root / ".codex_build" / "prez_zm"
out_dir.mkdir(parents=True, exist_ok=True)

summary = {}

wb = load_workbook(workbook_path, data_only=True)
summary["workbook_sheets"] = wb.sheetnames
workbook_rows = {}
for ws in wb.worksheets:
    rows = []
    for row in ws.iter_rows(values_only=True):
        vals = [v for v in row if v not in (None, "")]
        if vals:
            rows.append([str(v) for v in row])
    workbook_rows[ws.title] = rows[:80]
summary["workbook_rows"] = workbook_rows

doc = Document(docx_path)
paragraphs = []
for p in doc.paragraphs:
    text = " ".join(p.text.split())
    if text:
        paragraphs.append(text)
summary["docx_paragraph_count"] = len(paragraphs)
summary["docx_first_paragraphs"] = paragraphs[:120]

tables = []
for table in doc.tables:
    trows = []
    for row in table.rows:
        trows.append([" ".join(cell.text.split()) for cell in row.cells])
    tables.append(trows[:25])
summary["docx_tables"] = tables[:10]

with zipfile.ZipFile(pptm_path) as z:
    names = z.namelist()
    summary["pptm_slide_count"] = len([n for n in names if n.startswith("ppt/slides/slide") and n.endswith(".xml")])
    summary["pptm_has_macros"] = any(n == "ppt/vbaProject.bin" for n in names)
    summary["pptm_media"] = [n for n in names if n.startswith("ppt/media/")][:50]

(out_dir / "source_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps({
    "workbook_sheets": summary["workbook_sheets"],
    "docx_paragraph_count": summary["docx_paragraph_count"],
    "pptm_slide_count": summary["pptm_slide_count"],
    "pptm_has_macros": summary["pptm_has_macros"],
    "pptm_media_count": len(summary["pptm_media"]),
}, ensure_ascii=False, indent=2))
